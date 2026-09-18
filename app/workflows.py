"""The three dashboard actions: start this year's return, review a return, keying sheet for ATX.

Outputs are short and structured so the UI can show them as plain cards; nothing here needs
the language model, so they run in the offline profile and are the same on every model.
"""
from __future__ import annotations

import json
from typing import Any

from . import checklist, forecast
from .ingest.forms_catalog import BY_DOC_TYPE
from .ingest.transcript import FORM_TO_ITEM_KEY

MONEY_FACTS_1040 = ["wages", "taxable_interest", "ordinary_dividends", "capital_gain_loss", "schedule_c_net_profit", "adjusted_gross_income",
                    "taxable_income", "total_tax", "federal_withholding", "estimated_payments", "refund", "amount_owed"]
MONEY_FACTS_ENTITY = ["gross_receipts", "cost_of_goods_sold", "officer_compensation", "salaries_and_wages", "guaranteed_payments",
                      "ordinary_business_income", "distributions", "total_assets", "retained_earnings", "corporate_taxable_income", "corporate_total_tax"]
LABELS = {
    "wages": "Wages", "taxable_interest": "Interest", "ordinary_dividends": "Dividends", "capital_gain_loss": "Capital gains",
    "schedule_c_net_profit": "Business profit (Sch C)", "adjusted_gross_income": "Adjusted gross income", "taxable_income": "Taxable income",
    "total_tax": "Total tax", "federal_withholding": "Withholding", "estimated_payments": "Estimated payments", "refund": "Refund", "amount_owed": "Balance due",
    "gross_receipts": "Gross receipts", "cost_of_goods_sold": "Cost of goods sold", "officer_compensation": "Officer compensation",
    "salaries_and_wages": "Wages paid", "guaranteed_payments": "Guaranteed payments", "ordinary_business_income": "Ordinary income",
    "distributions": "Distributions", "total_assets": "Total assets", "retained_earnings": "Retained earnings",
    "corporate_taxable_income": "Taxable income", "corporate_total_tax": "Total tax", "underpayment_penalty": "Penalty",
}


def _label(name: str) -> str:
    return LABELS.get(name, name.replace("_", " ").capitalize())


def prior_return(store, client_id: str) -> dict[str, Any] | None:
    """The most recent return on file for the client, summarised for a card."""
    records = [r for r in store.canonical_records(client_id) if r.get("doc_type", "").startswith("form_") and r.get("tax_year")]
    if not records:
        return None
    r = sorted(records, key=lambda x: (x["tax_year"], x.get("return_type") != "none"))[-1]
    series = forecast.series_by_metric(store.facts_for_client(client_id))
    year = r["tax_year"]
    names = MONEY_FACTS_1040 if r.get("return_type") == "1040" else MONEY_FACTS_ENTITY
    figures = [{"label": _label(n), "name": n, "value": series[n][year]} for n in names if n in series and year in series[n]]
    return {"document_id": r["document_id"], "filename": r["filename"], "return_type": r.get("return_type"), "tax_year": year,
            "forms_present": r.get("forms_present", []), "filing_status": r.get("filing_status"), "figures": figures[:8],
            "notes": r.get("risk_flags", [])[:4], "summary": r.get("summary", "")}


def documents_this_year(store, client_id: str, tax_year: int) -> list[dict[str, Any]]:
    out = []
    for d in store.list_documents(client_id):
        if d.get("tax_year") == tax_year and d["status"] == "ready":
            spec = BY_DOC_TYPE.get(d.get("doc_type") or "")
            out.append({"id": d["id"], "filename": d["filename"], "doc_type": d.get("doc_type"), "form": spec.code if spec else None,
                        "plain": spec.name if spec else (d.get("doc_type") or "document").replace("_", " "), "uploaded_by": d.get("uploaded_by")})
    return out


def start_return(store, llm, client_id: str, tax_year: int | None, firm_name: str = "our office") -> dict[str, Any]:
    """Prep brief + request list for the new year. Idempotent: reuses an existing list for the year."""
    client = store.get_client(client_id)
    prior = prior_return(store, client_id)
    if tax_year is None:
        tax_year = (prior["tax_year"] + 1) if prior else None
    if tax_year is None:
        raise ValueError("Upload last year's return first so the archive knows what this client files.")
    existing = [l for l in store.list_request_lists(client_id) if l["tax_year"] == tax_year]
    if existing:
        rl = store.get_request_list(existing[0]["id"])
        created = False
    else:
        rl = checklist.build_for_client(store, llm, client_id, tax_year, firm_name)
        created = True
    pending = [i for i in rl["items"] if i["status"] == "pending"]
    received = [i for i in rl["items"] if i["status"] == "received"]
    risk = forecast.assess_risk(store.facts_for_client(client_id))
    watch = [f"{f['area']}: {f['detail']}" for f in risk["flags"][:3]]
    docs = documents_this_year(store, client_id, tax_year)
    steps = []
    if created or rl["status"] == "draft":
        steps.append("Read the request email, edit if needed, then mark it sent.")
    if pending:
        steps.append(f"Waiting on {len(pending)} item{'s' if len(pending) != 1 else ''} from the client.")
    if prior and prior.get("forms_present"):
        steps.append(f"Expect the same forms as last year: {', '.join(prior['forms_present'][:6])}{'…' if len(prior['forms_present']) > 6 else ''}.")
    if watch:
        steps.append("Review the watch items below before preparing.")
    return {"tax_year": tax_year, "client": client, "prior_return": prior, "request_list": {"id": rl["id"], "status": rl["status"], "pending": len(pending),
            "received": len(received), "total": len(rl["items"]), "created": created},
            "documents_received": docs, "watch": watch, "next_steps": steps,
            "headline": f"{client['name']}: {tax_year} {rl.get('return_type') or (prior or {}).get('return_type') or 'return'}. "
                        f"{len(received)} of {len(rl['items'])} items in, {len(pending)} still needed."}


def review_return(store, client_id: str, document_id: str) -> dict[str, Any]:
    """Compare a draft or filed return with the prior year's: what moved, what is missing, what to check."""
    doc = store.get_document(document_id)
    if not doc or doc["client_id"] != client_id:
        raise ValueError("That return is not in this client's folder.")
    canonical = json.loads(doc.get("canonical_json") or "{}")
    if not (doc.get("doc_type") or "").startswith("form_"):
        raise ValueError("That document is not a tax return. Upload the draft return (PDF from the tax software) and pick it.")
    year = doc.get("tax_year")
    series = forecast.series_by_metric(store.facts_for_client(client_id))
    prior_years = sorted({y for s in series.values() for y in s if year and y < year})
    prior_year = prior_years[-1] if prior_years else None
    changes = []
    for name, by_year in series.items():
        if year in by_year and prior_year in by_year:
            a, b = by_year[prior_year], by_year[year]
            if a == 0 and b == 0:
                continue
            pct = (b - a) / abs(a) if a else None
            changes.append({"label": _label(name), "name": name, "prior": a, "current": b, "change": round(b - a, 2), "pct": round(pct, 4) if pct is not None else None,
                            "big": pct is not None and abs(pct) >= 0.2 and abs(b - a) >= 1000})
    changes.sort(key=lambda c: (not c["big"], -abs(c["change"])))
    prior_forms: set[str] = set()
    for r in store.canonical_records(client_id):
        if r.get("tax_year") == prior_year and r.get("doc_type", "").startswith("form_"):
            prior_forms |= set(r.get("forms_present", []))
    current_forms = set(canonical.get("forms_present", []))
    missing = sorted(prior_forms - current_forms)
    new_forms = sorted(current_forms - prior_forms)
    risk = forecast.assess_risk(store.facts_for_client(client_id))
    checks = [f"{f['area']}: {f['detail']}" for f in risk["flags"] if risk.get("ratios", {}).get("year") == year][:5]
    only_current = [c for c in changes if c["big"]]
    verdict = []
    if missing:
        verdict.append(f"{len(missing)} form{'s' if len(missing) != 1 else ''} on last year's return not on this one: {', '.join(missing[:4])}.")
    if only_current:
        verdict.append(f"{len(only_current)} figure{'s' if len(only_current) != 1 else ''} moved more than 20%.")
    if checks:
        verdict.append(f"{len(checks)} planning item{'s' if len(checks) != 1 else ''} to check.")
    if not verdict:
        verdict.append("Nothing unusual against last year.")
    return {"document": {"id": doc["id"], "filename": doc["filename"], "return_type": canonical.get("return_type"), "tax_year": year},
            "prior_year": prior_year, "changes": changes[:15], "missing_forms": missing, "new_forms": new_forms, "checks": checks,
            "flags_from_return": canonical.get("risk_flags", []), "headline": " ".join(verdict)}


# ------------------------------------------------------------------ keying sheet for ATX
ATX_INPUT = {  # doc_type -> ATX input worksheet the preparer keys into
    "w2": "W-2 input", "1099_int": "1099-INT input (Sch B)", "1099_div": "1099-DIV input (Sch B)", "1099_b": "Form 8949 / Sch D detail (CSV import available)",
    "1099_da": "Form 8949 / Sch D detail (CSV import available)", "1099_nec": "Schedule C income", "1099_misc": "1099-MISC input", "1099_k": "Schedule C income",
    "1099_r": "1099-R input", "1099_g": "1099-G input", "ssa_1099": "SSA-1099 input", "1098_mortgage": "Schedule A mortgage interest",
    "1098_t": "Form 8863 / 1098-T input", "1098_e": "Student loan interest (Sch 1)", "1095_a": "Form 8962 / 1095-A input", "5498": "IRA contributions (Sch 1)",
    "1099_sa": "Form 8889", "5498_sa": "Form 8889", "k1_1065": "K-1 input worksheet (import from the 1065 if prepared in ATX)",
    "k1_1120s": "K-1 input worksheet (import from the 1120-S if prepared in ATX)", "k1_1041": "K-1 input worksheet",
    "trial_balance": "Trial balance import / 1120-S, 1065 or 1120 lines", "irs_transcript_wage_income": "Cross-check against inputs",
}


def keying_sheet(store, client_id: str, tax_year: int) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for d in store.list_documents(client_id):
        if d.get("tax_year") != tax_year or d["status"] != "ready":
            continue
        rec = json.loads(store.get_document(d["id"]).get("canonical_json") or "{}")
        target = ATX_INPUT.get(d.get("doc_type") or "")
        if rec.get("transcript_entries"):
            for e in rec["transcript_entries"]:
                rows.append({"atx_input": ATX_INPUT.get(_doc_type_for_form(e["form"]), "Cross-check"), "form": e["form"], "payer": e["payer"],
                             "field": e.get("amount_label") or "Amount", "value": e.get("amount"), "source": f"IRS transcript ({d['filename']})", "document_id": d["id"]})
            continue
        if not target:
            continue
        payer = next((e["name"] for e in rec.get("entities", []) if e.get("role") in ("employer", "payer", "lender", "issuer")), "")
        for f in rec.get("facts", []):
            rows.append({"atx_input": target, "form": (BY_DOC_TYPE.get(d.get("doc_type") or "").code if d.get("doc_type") in BY_DOC_TYPE else d.get("doc_type")),
                         "payer": payer, "field": _label(f["name"]), "value": f["value"], "source": f"{d['filename']}: {f.get('source_quote', '')}"[:160], "document_id": d["id"]})
    order = list(dict.fromkeys(ATX_INPUT.values()))
    rows.sort(key=lambda r: (order.index(r["atx_input"]) if r["atx_input"] in order else 99, r["form"] or "", r["payer"] or ""))
    return {"tax_year": tax_year, "rows": rows, "count": len(rows),
            "note": "Every figure with its source. Key into the ATX input worksheet named in the first column; Form 8949 detail can be imported from CSV."}


def _doc_type_for_form(form: str) -> str:
    key = FORM_TO_ITEM_KEY.get(form)
    mapping = {"W-2": "w2", "1099-INT": "1099_int", "1099-DIV": "1099_div", "1099-B": "1099_b", "1099-NEC": "1099_nec", "1099-MISC": "1099_misc",
               "1099-K": "1099_k", "1099-R": "1099_r", "1099-G": "1099_g", "SSA-1099": "ssa_1099", "1098": "1098_mortgage", "1098-T": "1098_t",
               "1098-E": "1098_e", "5498": "5498", "1099-SA": "1099_sa", "5498-SA": "5498_sa"}
    return mapping.get(form, key or "")


def keying_sheet_xlsx(sheet: dict[str, Any], client_name: str) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook(); ws = wb.active; ws.title = f"Keying {sheet['tax_year']}"
    ws.append([f"{client_name} — {sheet['tax_year']} keying sheet for ATX"]); ws["A1"].font = Font(bold=True)
    ws.append([sheet["note"]]); ws.append([])
    ws.append(["ATX input", "Form", "Payer", "Field", "Amount", "Source", "Keyed?"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in sheet["rows"]:
        ws.append([r["atx_input"], r["form"], r["payer"], r["field"], r["value"], r["source"], ""])
    for col, w in zip("ABCDEFG", (34, 12, 28, 26, 14, 60, 8)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=5):
        for c in row:
            c.number_format = "#,##0.00"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
